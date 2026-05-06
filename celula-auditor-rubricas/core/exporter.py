import os
import time
from openpyxl import Workbook, load_workbook

class ExcelExporter:
    """Módulo de exportación de resultados a Excel para SEA v4.0"""
    
    def __init__(self, output_path="Registro_Evaluaciones_PedagogIA.xlsx"):
        self.output_path = output_path
        self._initialize_file()

    def _initialize_file(self):
        """Crea el archivo con encabezados si no existe."""
        if not os.path.exists(self.output_path):
            wb = Workbook()
            ws = wb.active
            ws.title = "Evaluaciones"
            # Encabezados profesionales
            headers = [
                "Fecha/Hora", 
                "Archivo Alumno", 
                "Rúbrica Aplicada", 
                "Nota Final", 
                "Estado",
                "Criterios Evaluados"
            ]
            ws.append(headers)
            wb.save(self.output_path)

    def add_evaluation(self, student_file, rubric_name, nota, criteria_summary):
        """Añade una fila con el resultado de la evaluación."""
        try:
            wb = load_workbook(self.output_path)
            ws = wb.active
            
            estado = "Sobresaliente" if nota >= 9 else "Notable" if nota >= 7 else "Aprobado" if nota >= 5 else "Insuficiente"
            
            row = [
                time.strftime("%Y-%m-%d %H:%M:%S"),
                student_file,
                rubric_name,
                nota,
                estado,
                criteria_summary
            ]
            
            ws.append(row)
            wb.save(self.output_path)
            return True
        except Exception as e:
            print(f"Error al exportar a Excel: {e}")
            return False
