import os
import time
from openpyxl import Workbook, load_workbook

class ExcelExporter:
    """Módulo de exportación de resultados a Excel unificado."""
    
    def __init__(self, output_path="Registro_PedagogIA.xlsx"):
        self.output_path = output_path
        self._initialize_file()

    def _initialize_file(self):
        """Crea el archivo con encabezados estandarizados."""
        if not os.path.exists(self.output_path):
            wb = Workbook()
            ws = wb.active
            ws.title = "Resultados"
            headers = [
                "fecha_hora", 
                "entidad_id",
                "archivo_origen",
                "proceso_tipo", 
                "valor_resultado", 
                "nivel_estado",
                "detalles_resumen"
            ]
            ws.append(headers)
            wb.save(self.output_path)

    def add_row(self, entity_id, source_file, process_type, value, status_level, details):
        """Añade una fila con el resultado estandarizado."""
        try:
            wb = load_workbook(self.output_path)
            ws = wb.active
            
            row = [
                time.strftime("%Y-%m-%d %H:%M:%S"),
                entity_id,
                source_file,
                process_type,
                value,
                status_level,
                details
            ]
            
            ws.append(row)
            wb.save(self.output_path)
            return True
        except Exception as e:
            return False
